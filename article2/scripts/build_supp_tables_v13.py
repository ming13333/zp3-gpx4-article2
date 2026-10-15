# -*- coding: utf-8 -*-
"""
Update A2 supplementary tables for v13 (TO submission).

 - S5: add the "on-treatment only" tier and the mean-aggregation patient-level rule
 - S6: add the "On_only" framework for the 11-gene benchmark panel
 - S9: NEW table of patient-cluster permutation null distributions (four frameworks)

Reads the frozen framework masks + FPKM matrix and results/a2_v13_cluster_null_frozen.json.
Writes results/A2_Supplementary_Tables_v13.xlsx (a copy; the v9 workbook is untouched).
"""
import csv, json, os, sys, shutil
import numpy as np
from scipy import stats
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

sys.stdout.reconfigure(encoding="utf-8")
# --- portable paths: everything resolves relative to this file's repository ---
_HERE = os.path.dirname(os.path.abspath(__file__))   # <repo>/article2/scripts
A2    = os.path.dirname(_HERE)                       # <repo>/article2
ROOT  = os.path.dirname(A2)                          # <repo>
FPKM  = os.environ.get("A2_GSE91061_FPKM", os.path.join(
        ROOT, "output", "gse91061_validation",
        "GSE91061_BMS038109Sample.hg19KnownGene.fpkm.csv"))
MASKS = os.path.join(A2, "results", "framework_cohort_masks_v2.json")
SRC   = os.path.join(A2, "results", "A2_Supplementary_Tables.xlsx")
DST   = os.path.join(A2, "results", "A2_Supplementary_Tables_v13.xlsx")
NULLJ = os.path.join(A2, "results", "a2_v13_cluster_null_frozen.json")
if not os.path.exists(FPKM):
    raise SystemExit(
        "Input FPKM matrix not found: %s\n"
        "Obtain GSE91061 from GEO and place "
        "GSE91061_BMS038109Sample.hg19KnownGene.fpkm.csv at that path, or set A2_GSE91061_FPKM." % FPKM)

# ---------- load ----------
with open(FPKM, newline="", encoding="utf-8") as f:
    r = csv.reader(f); hdr = next(r)
    samples_all = [h.strip() for h in hdr[1:]]
    ids, rows = [], []
    for row in r:
        if not row:
            continue
        ids.append(row[0].strip().strip('"'))
        rows.append([float(x) for x in row[1:]])
X = np.array(rows, float)
sid = {s: i for i, s in enumerate(samples_all)}
M = json.load(open(MASKS, encoding="utf-8"))
lab = M["labels"]
A = M["A_pre"]; B = M["B_naive_pooled"]; C = M["C_patient_pre_else_on"]
ON = [s for s in B if "_On_" in s]

PANEL = [("CD8A", 925), ("GZMB", 3002), ("PRF1", 5551), ("PDCD1", 5133), ("CD274", 29126),
         ("CTLA4", 1493), ("LAG3", 3902), ("TIGIT", 201633), ("IFNG", 3458),
         ("CXCL9", 4283), ("ZP3", 7784)]
ENTREZ = {e: n for n, e in PANEL}
# row identifiers in the FPKM matrix are Entrez IDs (as strings)
idx_of = {}
for i, g in enumerate(ids):
    gg = g.strip().strip('"')
    idx_of[gg] = i
    try:
        idx_of[int(gg)] = i
    except ValueError:
        pass


def mwu(vals_r, vals_n):
    u, p = stats.mannwhitneyu(vals_r, vals_n, alternative="two-sided")
    n1, n2 = len(vals_r), len(vals_n)
    rrb = 2 * u / (n1 * n2) - 1
    se = np.sqrt((n1 + n2 + 1) / (3 * n1 * n2))
    return u, p, rrb, rrb - 1.96 * se, rrb + 1.96 * se


def n_patients(cols):
    return len(set(c.split("_")[0] for c in cols))


def boot_ci(cols, gene, B=20000, seed=20261014):
    """stratified bootstrap CI (resample within response groups), as in the frozen pipeline"""
    gi = idx_of[gene]
    ia = [sid[s] for s in cols if lab[s] == "R"]
    ib = [sid[s] for s in cols if lab[s] == "N"]
    vr = X[gi, ia]; vn = X[gi, ib]
    rng = np.random.default_rng(seed)
    out = np.empty(B)
    for i in range(B):
        a = vr[rng.integers(0, len(vr), len(vr))]
        b = vn[rng.integers(0, len(vn), len(vn))]
        u, _ = stats.mannwhitneyu(a, b, alternative="two-sided")
        out[i] = 2 * u / (len(a) * len(b)) - 1
    return float(np.percentile(out, 2.5)), float(np.percentile(out, 97.5))


def frame_stats(cols, gene):
    gi = idx_of[gene]
    ia = [sid[s] for s in cols if lab[s] == "R"]
    ib = [sid[s] for s in cols if lab[s] == "N"]
    vr = X[gi, ia]; vn = X[gi, ib]
    u, p, rrb, lo, hi = mwu(vr, vn)
    return dict(n=len(cols), npt=n_patients(cols), nR=len(ia), nN=len(ib),
                med_r=float(np.median(vr)), med_n=float(np.median(vn)),
                U=float(u), P=float(p), r=float(rrb), lo=float(lo), hi=float(hi))


# ---------- mean-aggregation framework ----------
def mean_agg(cols):
    """one value per patient = mean of that patient's included samples"""
    byp = {}
    for c in cols:
        byp.setdefault(c.split("_")[0], []).append(sid[c])
    out_r, out_n, npat = {}, {}, 0
    for pid, ix in byp.items():
        v = X[:, ix].mean(axis=1)
        if lab[[c for c in cols if c.split("_")[0] == pid][0]] == "R":
            out_r[pid] = v
        else:
            out_n[pid] = v
        npat += 1
    return out_r, out_n, npat


mr, mn, npat = mean_agg(B)


def mean_agg_stats(gene):
    gi = idx_of[gene]
    vr = np.array([v[gi] for v in mr.values()])
    vn = np.array([v[gi] for v in mn.values()])
    u, p, rrb, lo, hi = mwu(vr, vn)
    return dict(n=npat, nR=len(vr), nN=len(vn),
                med_r=float(np.median(vr)), med_n=float(np.median(vn)),
                U=float(u), P=float(p), r=float(rrb), lo=float(lo), hi=float(hi))


print("On_only ZP3:", frame_stats(ON, 7784))
print("mean-agg ZP3:", mean_agg_stats(7784))

# ---------- open workbook copy ----------
shutil.copyfile(SRC, DST)
wb = openpyxl.load_workbook(DST)
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
HDR = PatternFill("solid", fgColor="DDEBF7")

# ================= S5 rebuild =================
ws = wb["S5"]
for r in range(3, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        ws.cell(row=r, column=c).value = None

st_pre = frame_stats(A, 7784)
st_on = frame_stats(ON, 7784)
st_all = frame_stats(B, 7784)
st_c = frame_stats(C, 7784)
st_m = mean_agg_stats(7784)

rows5 = [
    ("Pre (baseline) only", st_pre),
    ("On-treatment only", st_on),
    ("All samples (incl. on-treatment)", st_all),
    ("Patient-level, pre-else-on rule", st_c),
    ("Patient-level, mean-aggregation rule", st_m),
]
ws.cell(row=3, column=1, value="Scope").font = BOLD
ws.cell(row=3, column=2, value="n samples").font = BOLD
ws.cell(row=3, column=3, value="n patients").font = BOLD
ws.cell(row=3, column=4, value="n responders").font = BOLD
ws.cell(row=3, column=5, value="n non-responders").font = BOLD
ws.cell(row=3, column=6, value="Median ZP3 (resp)").font = BOLD
ws.cell(row=3, column=7, value="Median ZP3 (non)").font = BOLD
ws.cell(row=3, column=8, value="U").font = BOLD
ws.cell(row=3, column=9, value="P").font = BOLD
ws.cell(row=3, column=10, value="r_rb").font = BOLD
ws.cell(row=3, column=11, value="CI low").font = BOLD
ws.cell(row=3, column=12, value="CI high").font = BOLD
for i, (name, s) in enumerate(rows5):
    r = 4 + i
    ws.cell(row=r, column=1, value=name)
    ws.cell(row=r, column=2, value=s["n"])
    ws.cell(row=r, column=3, value=s.get("npt", s["n"]))
    ws.cell(row=r, column=4, value=s["nR"])
    ws.cell(row=r, column=5, value=s["nN"])
    ws.cell(row=r, column=6, value=round(s["med_r"], 4))
    ws.cell(row=r, column=7, value=round(s["med_n"], 4))
    ws.cell(row=r, column=8, value=s["U"])
    ws.cell(row=r, column=9, value=round(s["P"], 6))
    ws.cell(row=r, column=10, value=round(s["r"], 4))
    ws.cell(row=r, column=11, value=round(s["lo"], 4))
    ws.cell(row=r, column=12, value=round(s["hi"], 4))
ws.cell(row=10, column=1,
        value=("Footnote: Mann-Whitney U test of ZP3 by response in GSE91061 under five scopes. "
               "'On-treatment only' uses the 56 on-treatment samples from 55 patients and is the reference for the amount of genuine "
               "treatment-induced signal available in this cohort. Patient-level aggregation is reported under two rules because the rule is "
               "not dictated by the data: 'pre-else-on' takes the pre-treatment sample when available and otherwise the on-treatment sample; "
               "'mean-aggregation' takes the mean of each patient's available samples. Neither rule reaches P<0.05, but the estimates differ "
               "(P=0.175 vs P=0.054), so the result should not be described as absent from every framework. "
               "Approximate 95% CIs use r_rb +/- 1.96*sqrt((n1+n2+1)/(3*n1*n2)). "
               "Source: results/framework_cohort_masks_v2.json; scripts/a2_v13_cluster_null_frozen.py; "
               "GSE91061_BMS038109Sample.hg19KnownGene.fpkm.csv (Entrez 7784 = ZP3)."))
ws.cell(row=10, column=1).alignment = WRAP

# ================= S6 append On_only =================
ws6 = wb["S6"]
start = ws6.max_row + 1
# locate column order from header row 3
hdr6 = [ws6.cell(row=3, column=c).value for c in range(1, 13)]
print("S6 header:", hdr6)
r = start
for gene, ez in PANEL:
    s = frame_stats(ON, ez)
    ws6.cell(row=r, column=1, value=gene)
    ws6.cell(row=r, column=2, value=ez)
    ws6.cell(row=r, column=3, value="On_only")
    ws6.cell(row=r, column=4, value=s["nR"])
    ws6.cell(row=r, column=5, value=s["nN"])
    ws6.cell(row=r, column=6, value=round(s["med_r"], 3))
    ws6.cell(row=r, column=7, value=round(s["med_n"], 3))
    ws6.cell(row=r, column=8, value=round(s["P"], 4))
    ws6.cell(row=r, column=9, value=round(s["r"], 3))
    ws6.cell(row=r, column=10, value=f"[{s['lo']:.3f}, {s['hi']:.3f}]")
    blo, bhi = boot_ci(ON, ez)
    ws6.cell(row=r, column=11, value=f"[{blo:.3f}, {bhi:.3f}]")
    ws6.cell(row=r, column=12, value=bool(s["P"] < 0.05))
    r += 1
ws6.cell(row=r + 1, column=1,
         value=("Footnote (v13 addition): the On_only framework (56 on-treatment samples from 55 patients) was added so that the benchmark panel "
                "spans the full range of frameworks. Because the on-treatment markers genuinely rise on treatment (see Table S8), their nominal "
                "significance under naive pooling (Framework B) reflects real treatment-induced biology rather than a false positive; the point of "
                "the panel is that a pooled analysis cannot distinguish a baseline-predictive association from a pharmacodynamic one."))
ws6.cell(row=r + 1, column=1).alignment = WRAP

# ================= S9 new =================
if "S9" in wb.sheetnames:
    del wb["S9"]
ws9 = wb.create_sheet("S9")
N = json.load(open(NULLJ, encoding="utf-8"))
PERM = N["perm"]; OBS = N["frozen_caliber_observed"]
LAB = {"A_pre": "A - Pre-treatment only",
       "On_only": "On-treatment only",
       "B_naive": "B - Naive pooled (pre + on)",
       "C_patient": "C - Patient-level aggregated (pre-else-on)"}
ws9.cell(row=1, column=1, value="Table S9. Patient-cluster permutation null distributions for the GSE91061 transcriptome-wide audit (22,187 genes)").font = Font(bold=True, size=12)
ws9.cell(row=2, column=1, value=("Null distributions were built from 500 permutations of patient labels; the sample-level Mann-Whitney U statistic was recomputed on each permutation, "
                                 "preserving the within-patient repeated-measure structure and the number of responder patients. Per-gene P values use the tie-corrected, continuity-corrected "
                                 "normal approximation, which reproduces the frozen per-gene scipy.stats.mannwhitneyu counts exactly (832 / 4,015 / 4,334 / 843 at P<0.05; 0.0% deviation). "
                                 "The 'null share' column is the null distribution mean expressed as a percentage of the nominal 5% threshold: a framework above 5% is anti-conservative."))
ws9.cell(row=2, column=1).alignment = WRAP
ws9.row_dimensions[2].height = 48
cols9 = ["Framework", "n samples", "n patients", "n responder patients",
         "Observed genes P<0.05 (scipy)", "Observed genes P<0.05 (vectorised)",
         "Null mean (genes)", "Null p95 (genes)", "Null share of 5% threshold (%)",
         "Permutation P", "Excess over null"]
for j, c in enumerate(cols9, start=1):
    cell = ws9.cell(row=4, column=j, value=c)
    cell.font = BOLD; cell.fill = HDR; cell.alignment = WRAP
order = ["A_pre", "On_only", "B_naive", "C_patient"]
for i, k in enumerate(order):
    p = PERM[k]
    vals = [LAB[k], p["n"], p["n_patients"], p["nR_patients"],
            p["observed_scipy"], p["observed_vectorized"],
            round(p["null_mean"]), round(p["null_p95"]), p["null_pct"],
            round(p["perm_P"], 4), p["excess"]]
    for j, v in enumerate(vals, start=1):
        ws9.cell(row=5 + i, column=j, value=v)
ws9.cell(row=10, column=1,
         value=("Reading: under a true null the same sample-level test occupies 4.6-4.9% of the nominal 5% threshold in Frameworks A, On_only and C, but 10.1% under naive pooling (B) - "
                "i.e. pooling approximately doubles the false-positive floor. The rise in the observed count from 832 to 4,334 is nevertheless dominated by genuine on-treatment biology: "
                "the On_only framework alone yields 4,015 genes on a valid 4.9% null floor (permutation P=0.002; excess +2,927). Pooling therefore raises the false-positive floor AND "
                "dilutes the real signal (excess +2,927 -> +2,092). Neither A (excess -185; P=0.558) nor C (excess -191; P=0.614) shows any excess over its null. "
                "Resolution note: with 500 permutations the smallest attainable permutation P for the transcriptome-wide analyses is 0.002."))
ws9.cell(row=10, column=1).alignment = WRAP
for col, w in zip("ABCDEFGHIJK", [40, 10, 11, 12, 16, 16, 12, 12, 14, 12, 12]):
    ws9.column_dimensions[col].width = w

wb.save(DST)
print("SAVED", DST)
print("sheets:", wb.sheetnames)
